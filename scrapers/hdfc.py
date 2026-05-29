"""HDFC AMC scraper — downloads per-fund XLSX portfolio disclosures."""
import re
import io
import datetime
import requests
import openpyxl
from .base import parse_date_from_text, clean_str, safe_float

DISCLOSURE_URL = "https://www.hdfcfund.com/investor-service/portfolio-disclosure"
BASE_URL = "https://www.hdfcfund.com"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Referer": "https://www.hdfcfund.com/",
}

# Filename keyword → scheme_code
FUND_NAME_TO_SCHEME = {
    "flexi cap":                        "118955",
    "small cap":                        "130503",
    "mid cap":                          "118989",
    "large cap":                        "119018",
    "focused":                          "118950",
    "multi cap":                        "149368",
    "elss tax saver":                   "119060",
    "elss tax":                         "119060",
    "hybrid equity":                    "119062",
    "hybrid debt":                      "119118",
    "multi-asset":                      "119131",
    "arbitrage":                        "118931",
    "corporate bond":                   "118987",
    "credit risk":                      "128051",
    "short term  debt":                 "119016",
    "short term debt":                  "119016",
    "medium term":                      "119081",
    "dynamic debt":                     "119075",
    "income fund":                      "119069",
    "gilt fund":                        "119116",
    "liquid fund":                      "119091",
    "money market":                     "119092",
    "overnight":                        "119110",
    "ultra short term":                 "145034",
    "long duration":                    "151313",
    "bse 500 index":                    "151728",
    "bse india sector leaders":         "153959",
    "crisil-ibx financial services 3-6":"153517",
    "crisil-ibx financial services 9-12":"154308",
    "nifty 100 index":                  "149868",
    "nifty 100 equal weight":           "149870",
    "nifty50 equal weight":             "149107",
    "nifty next 50":                    "149288",
    "nifty midcap 150":                 "151724",
    "nifty smallcap 250":               "151727",
    "nifty largemidcap 250":            "152889",
    "nifty g-sec jun 2027":             "151181",
    "nifty g-sec dec 2026":             "150845",
    "nifty g-sec apr 2029":             "151495",
    "nifty g-sec jul 2031":             "150847",
    "nifty g-sec sep 2032":             "151183",
    "nifty g-sec jun 2036":             "151489",
    "nifty sdl oct 2026":               "151456",
    "nifty sdl plus g-sec jun 2027":    "151570",
    "nifty india consumption":          "154179",
    "nifty india digital":              "153097",
    "developed world":                  "149180",
}


def _parse_sheet(ws) -> dict:
    """
    HDFC layout (0-indexed):
      col 0: row marker ('|' or empty)
      col 1: ISIN
      col 2: coupon (empty for equity)
      col 3: security name
      col 4: sector / rating
      col 5: quantity
      col 6: market value (Rs. in Lacs)
      col 7: % to NAV
    """
    date_str = None
    holdings = []

    for row in ws.iter_rows(values_only=True):
        cells = [clean_str(c) for c in row]
        text = " ".join(cells)

        if not date_str:
            # "Portfolio as on 30-Apr-2026"
            m = re.search(r'portfolio as on\s+(\d{1,2}-\w{3}-\d{4})', text, re.IGNORECASE)
            if m:
                try:
                    date_str = datetime.datetime.strptime(m.group(1), "%d-%b-%Y").strftime("%Y-%m-%d")
                except ValueError:
                    pass
            if not date_str and "as on" in text.lower():
                date_str = parse_date_from_text(text)

        if len(cells) < 8:
            continue

        isin   = cells[1]
        name   = cells[3]
        sector = cells[4]

        if not isin or not isin.startswith("IN"):
            continue
        if not name:
            continue
        if re.match(r"^(sub\s*total|total|grand\s*total)", name.strip(), re.IGNORECASE):
            continue

        quantity = safe_float(cells[5])
        mkt_val  = safe_float(cells[6])
        pct_nav  = safe_float(cells[7])

        holdings.append({
            "isin":              isin,
            "name":              name,
            "sector":            sector,
            "quantity":          quantity,
            "market_value_lakh": mkt_val,
            "pct_nav":           pct_nav,
        })

    return {"date": date_str, "holdings": holdings}


def fetch() -> dict:
    """Fetch latest HDFC portfolio disclosures."""
    session = requests.Session()
    session.headers.update(HEADERS)

    resp = session.get(DISCLOSURE_URL, timeout=30)
    resp.raise_for_status()

    pattern = r'href=["\']([^"\']*Monthly[^"\']*\.xlsx[^"\']*)["\']'
    links = re.findall(pattern, resp.text, re.IGNORECASE)
    xlsx_links = [l for l in links if "2026" in l or "2025" in l]

    if not xlsx_links:
        raise RuntimeError("Could not find HDFC XLSX links on disclosure page")

    print(f"  HDFC: found {len(xlsx_links)} files")

    schemes = []
    global_date = None

    for path in xlsx_links:
        url = (BASE_URL + path) if path.startswith("/") else path
        filename = url.split("/")[-1].split("?")[0].lower()

        scheme_code = None
        for kw, code in FUND_NAME_TO_SCHEME.items():
            if kw in filename:
                scheme_code = code
                break

        print(f"  HDFC: downloading {filename[:60]}")
        try:
            dl = session.get(url, timeout=60)
            dl.raise_for_status()
        except Exception as e:
            print(f"  HDFC: WARNING — failed: {e}")
            continue

        wb = openpyxl.load_workbook(io.BytesIO(dl.content), data_only=True)
        # Use first non-derivative sheet
        ws = next(
            (wb[s] for s in wb.sheetnames if "derivative" not in s.lower()),
            wb.active
        )
        result = _parse_sheet(ws)

        if result["date"] and not global_date:
            global_date = result["date"]

        if not result["holdings"]:
            print(f"  HDFC: WARNING — no holdings in {filename[:40]}")
            continue

        schemes.append({
            "sheet_code":  wb.sheetnames[0],
            "scheme_code": scheme_code,
            "filename":    filename,
            "holdings":    result["holdings"],
        })

    return {
        "amc":     "hdfc",
        "date":    global_date,
        "source":  DISCLOSURE_URL,
        "schemes": schemes,
    }
