"""Kotak AMC scraper — downloads consolidated SEBI portfolio XLSX (all funds, one file)."""
import re
import io
import requests
import openpyxl
from .base import parse_date_from_text, clean_str, safe_float

DISCLOSURE_URL = "https://www.kotakamc.com/downloads"
BASE_URL = "https://www.kotakamc.com"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Referer": "https://www.kotakamc.com/",
}

# Sheet code → scheme_code for funds in our universe
SHEET_TO_SCHEME = {
    "SEF": "120166",  # Kotak Flexicap Fund
    "NEF": "119775",  # Kotak Midcap Fund
    "KBC": "120152",  # Kotak Large Cap Fund
    "MSC": "120164",  # Kotak Small Cap Fund
    "BAL": "133035",  # Kotak Aggressive Hybrid Fund
    "NIF": "148978",  # Kotak Nifty 50 Index Fund
    "MCF": "149185",  # Kotak Multicap Fund
    "MID": "120158",  # Kotak Large & Midcap Fund
    "ELS": "119773",  # Kotak ELSS Tax Saver Fund
}


def _parse_sheet(ws) -> dict:
    """
    Kotak layout (0-indexed):
      col 0: section header
      col 1: sub-type / space
      col 2: security name
      col 3: ISIN
      col 4: sector / rating
      col 5: yield
      col 6: quantity
      col 7: market value (Rs. in Lacs)
      col 8: % to Net Assets
    """
    date_str = None
    holdings = []

    for row in ws.iter_rows(values_only=True):
        cells = [clean_str(c) for c in row]
        text = " ".join(cells)

        if not date_str and "as on" in text.lower():
            date_str = parse_date_from_text(text)

        if len(cells) < 9:
            continue

        name   = cells[2]
        isin   = cells[3]
        sector = cells[4]

        if not name or not isin or not isin.startswith("IN"):
            continue
        if re.match(r"^(sub\s*total|total|grand\s*total)", name.strip(), re.IGNORECASE):
            continue

        quantity = safe_float(cells[6])
        mkt_val  = safe_float(cells[7])
        pct_nav  = safe_float(cells[8])

        holdings.append({
            "isin":              isin,
            "name":              name,
            "sector":            sector,
            "quantity":          quantity,
            "market_value_lakh": mkt_val,
            "pct_nav":           pct_nav,
        })

    return {"date": date_str, "holdings": holdings}


def fetch() -> dict:
    """Fetch latest Kotak consolidated portfolio XLSX."""
    session = requests.Session()
    session.headers.update(HEADERS)

    resp = session.get(DISCLOSURE_URL, timeout=30)
    resp.raise_for_status()

    # Find consolidated XLSX link — pattern: ConsolidatedSEBIPortfolio<Month><Year>.xlsx
    pattern = r'href=["\']([^"\']*ConsolidatedSEBI[^"\']*\.xlsx[^"\']*)["\']'
    matches = re.findall(pattern, resp.text, re.IGNORECASE)
    if not matches:
        raise RuntimeError("Could not find Kotak consolidated XLSX link on downloads page")

    url = matches[0]
    if url.startswith("/"):
        url = BASE_URL + url

    print(f"  Kotak: downloading {url.split('?')[0].split('/')[-1]}")
    dl = session.get(url, timeout=120)
    dl.raise_for_status()

    wb = openpyxl.load_workbook(io.BytesIO(dl.content), data_only=True)

    schemes = []
    global_date = None

    for sheet_code, scheme_code in SHEET_TO_SCHEME.items():
        if sheet_code not in wb.sheetnames:
            print(f"  Kotak: sheet {sheet_code} not found, skipping")
            continue

        ws = wb[sheet_code]
        result = _parse_sheet(ws)

        if result["date"] and not global_date:
            global_date = result["date"]

        if not result["holdings"]:
            print(f"  Kotak: WARNING — no holdings in {sheet_code}")
            continue

        schemes.append({
            "sheet_code":  sheet_code,
            "scheme_code": scheme_code,
            "holdings":    result["holdings"],
        })

    return {
        "amc":     "kotak",
        "date":    global_date,
        "source":  DISCLOSURE_URL,
        "schemes": schemes,
    }
